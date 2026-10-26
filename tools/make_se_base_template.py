"""Derive the report base template from the SE Deck Kit template.

The SE Deck Kit ships SE_Deck_Template.pptx with three layouts (COVER,
CONTENT, SECTION), the SE theme (Cambria / Calibri, SE colours), the logos and
21 guide / sample slides.  The report generator only needs the layouts and the
theme, so this script removes every slide and the section list and writes
static/se_deck/SE_Deck_Base.pptx.  Re-run it when the SE template changes:

    python tools/make_se_base_template.py path/to/SE_Deck_Template.pptx
"""
import os
import sys

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "static", "se_deck", "SE_Deck_Base.pptx")
NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def main(src: str, out: str = OUT) -> None:
    prs = Presentation(src)
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)          # the slide part (and its notes) is no longer saved
        sld_ids.remove(sld_id)
    # The section list names the removed slides; PowerPoint asks to repair a file
    # whose sections point at missing slides, so drop it (the generator writes its own).
    ext_lst = prs.part._element.find(f"{NS_P}extLst")
    if ext_lst is not None:
        for ext in list(ext_lst):
            if "sectionLst" in "".join(child.tag for child in ext):
                ext_lst.remove(ext)
        if len(ext_lst) == 0:
            prs.part._element.remove(ext_lst)
    names = [layout.name for layout in prs.slide_layouts]
    for need in ("COVER", "CONTENT", "SECTION"):
        if need not in names:
            raise SystemExit(f"layout {need!r} not found in {src} (found {names})")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print(f"wrote {out}: layouts {names}, {len(prs.slides)} slides")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    main(sys.argv[1], *(sys.argv[2:3]))
