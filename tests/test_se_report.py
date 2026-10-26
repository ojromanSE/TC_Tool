"""Tests for se_report.py — the SE-standard PPTX / PDF type-curve report."""
from __future__ import annotations

import datetime as dt
import re
import shutil
import subprocess
import zipfile

import numpy as np
import pandas as pd
import pytest

import core
import se_report as R

D_LIM = core.D_LIM_DEFAULT


# ---------------------------------------------------------------- helpers
def _oneline(n=30, seed=0, eur_col="EUR (Mbbl)", scale=1.0):
    rng = np.random.default_rng(seed)
    eur = rng.lognormal(np.log(450 * scale), 0.35, n)
    eur[0] *= 6.0                                   # one high outlier
    return pd.DataFrame({
        "WellID": [f"W{i:02d}_1H" for i in range(n)],
        "WellName": [f"TESTWELL {i:02d}" for i in range(n)],
        "API10": [f"42{i:08d}" for i in range(n)],
        "State": "TX", "County": ["MARTIN" if i % 3 else "HOWARD" for i in range(n)],
        "PrimaryFormation": ["WOLFCAMP A" if i % 2 else "WOLFCAMP B" for i in range(n)],
        "LateralLength": rng.uniform(6000, 11000, n).round(),
        "CompletionDate": pd.date_range("2017-01-01", periods=n, freq="45D"),
        "FirstProdDate": pd.date_range("2017-02-01", periods=n, freq="45D"),
        "Months History": rng.integers(18, 60, n),
        "Model": "Modified Arps",
        "qi (per day)": rng.uniform(300, 900, n) * scale,
        "b": rng.uniform(0.8, 1.2, n).round(3),
        "di (per month)": rng.uniform(0.06, 0.2, n),
        "b at bound": [i % 7 == 0 for i in range(n)],
        eur_col: eur,
    })


def _curves(f: R.FluidReport):
    t = np.arange(1, 601, dtype=float)
    data = {"t": t}
    for k, g in f.groups.items():
        q = g["qi"] * core.DAYS_PER_MONTH
        data[k] = (core.modified_arps_cum(q, g["b"], g["di"], D_LIM, t)
                   - core.modified_arps_cum(q, g["b"], g["di"], D_LIM, t - 1))
    lines = [(np.arange(1, 25, dtype=float), 8000 * np.exp(-0.08 * np.arange(24)), "W01_1H")]
    return pd.DataFrame(data), lines


def _fluid(name="Oil", eur_col="EUR (Mbbl)", n=30, seed=0, exclude=(0,), scale=1.0):
    ol = _oneline(n, seed, eur_col, scale)
    sel = set(ol["WellID"]) - {ol["WellID"].iloc[i] for i in exclude}
    f = R.fluid_report(name, ol, sel, eur_col, None, None, D_LIM, 10000, True, 0.8, 1.2)
    f.curves, f.lines = _curves(f)
    return f, ol


def _input(fluids, name="Wolfcamp Martin 2017-21", disclaimer="Read this. " * 20):
    an = R.analog_summary(fluids, fluids[0].wells.assign(State="TX", PrimaryFormation="WOLFCAMP A",
                                                         FirstProdDate="2017-03-01"))
    return R.ReportInput(
        tc_name=name, source="Production data", as_of="Sep 2025", generated=dt.date(2026, 10, 26),
        counties=an["counties"], state="TX", formations=["WOLFCAMP A", "WOLFCAMP B"], vintage="2017–21",
        lateral_min=an["lateral_min"], lateral_max=an["lateral_max"], lateral_mean=an["lateral_mean"],
        months_median=an["months_median"], norm_len=10000, use_norm=True, exponent=1.0,
        settings=[("Normalization length", "10,000 ft"), ("b-factor range", "0.80–1.20")],
        decline_models="Modified Arps and SEPD", fluids=list(fluids), disclaimer=disclaimer)


@pytest.fixture(scope="module")
def three_phase():
    oil, _ = _fluid("Oil", "EUR (Mbbl)", seed=0)
    gas, _ = _fluid("Gas", "EUR (MMcf)", seed=1, exclude=(), scale=2.4)
    wtr, _ = _fluid("Water", "EUR (Mbbl water)", seed=2, exclude=(0, 1), scale=0.5)
    return _input([oil, gas, wtr])


@pytest.fixture(scope="module")
def rendered(three_phase, tmp_path_factory):
    deck = R.build_deck(three_phase)
    d = tmp_path_factory.mktemp("deck")
    pptx = R.render_pptx(deck, str(d / (deck.file_stem + ".pptx")))
    pdf = R.render_pdf(deck, str(d / (deck.file_stem + ".pdf")))
    return deck, pptx, pdf


# ---------------------------------------------------------------- text and numbers (SE section 4)
def test_number_formats():
    assert R.fmt_vol(1172.4) == "1,172"
    assert R.fmt_vol(48.44) == "48.4"
    assert R.fmt_vol(13.0) == "13"                      # '.0' dropped
    assert R.fmt_vol(0.537) == "0.54"
    assert R.fmt_vol(float("nan")) == "–"
    assert R.fmt_int(10000) == "10,000"
    assert R.fmt_pct(0.574) == "57%"
    assert R.fmt_ratio(1.63) == "1.6"


def test_nb_keeps_numbers_with_units():
    s = R.nb("484 Mbbl at 10,000 ft and 26 wells")
    assert "484 Mbbl" in s and "10,000 ft" in s and "26 wells" in s


def test_title_name_and_join():
    assert R.title_name("BLUE MESA 12") == "Blue Mesa 12"
    assert R.title_name("Wolfcamp A") == "Wolfcamp A"
    assert R.join_and(["Oil", "Gas", "Water"]) == "Oil, Gas and Water"
    assert R.join_and(["Oil"]) == "Oil"


def test_formation_phrase_merges_shared_names(three_phase):
    assert R._formation_phrase(three_phase) == "Wolfcamp A and B"


# ---------------------------------------------------------------- report data
def test_fluid_report_type_wells_reproduce_their_eur():
    f, ol = _fluid()
    assert f.n_total == 30 and f.n_tc == 29
    for key in ("P90", "P50", "P10", "Mean"):
        g = f.groups[key]
        cum = core.modified_arps_cum(g["qi"] * core.DAYS_PER_MONTH, g["b"], g["di"], D_LIM, np.array([600.0]))[0]
        assert cum / 1000.0 == pytest.approx(f.stats[key], rel=1e-9)      # curve volume == EUR target
        q0 = core.modified_arps(1.0, g["b"], g["di"], D_LIM, np.array([0.0]))[0]
        q12 = core.modified_arps(1.0, g["b"], g["di"], D_LIM, np.array([12.0]))[0]
        assert g["dec1"] == pytest.approx(1 - q12 / q0)
    assert f.groups["Mean"]["n"] == 29
    assert f.groups["P10"]["n"] == 2 and f.groups["P90"]["n"] == 3
    assert f.per_ft(f.stats["P50"]) == pytest.approx(f.stats["P50"] * 1000 / 10000)
    assert f.n_b_at_bound == int(ol.loc[ol["WellID"] != "W00_1H", "b at bound"].sum())
    assert bool(f.wells.loc[f.wells["WellID"] == "W00_1H", "in_tc"].iloc[0]) is False


def test_fluid_report_handles_missing_data():
    assert R.fluid_report("Oil", pd.DataFrame(), None, "EUR (Mbbl)", None, None, D_LIM, 10000, True, 0.8, 1.2) is None
    ol = _oneline(8)
    f = R.fluid_report("Oil", ol, set(), "EUR (Mbbl)", None, None, D_LIM, 10000, True, 0.8, 1.2)
    assert f.n_tc == 0 and f.groups == {}


# ---------------------------------------------------------------- deck content
def test_deck_flow_and_sections(three_phase):
    deck = R.build_deck(three_phase)
    titles = [s.fields.get("title") for s in deck.slides]
    assert deck.slides[0].layout == "COVER" and titles[0] == "Wolfcamp Martin 2017-21 Type Curve"
    assert titles[1:3] == ["Disclaimer", "EUR Summary by Phase"]
    assert deck.slides[1].fields["sub"] == ""                    # disclaimer has no subtitle
    assert "Type Curve at a Glance" not in titles and "Analog Set and Method" not in titles
    for fl in ("Oil", "Gas", "Water"):
        assert [f"{fl} Type Curve", f"{fl} EUR Distribution and Decline Parameters"] == \
               [t for t in titles if t and t.startswith(fl)]
        assert not any(t and t.startswith("Outlier") for t in titles)
    assert titles[3:9] == ["Oil Type Curve", "Oil EUR Distribution and Decline Parameters",
                           "Gas Type Curve", "Gas EUR Distribution and Decline Parameters",
                           "Water Type Curve", "Water EUR Distribution and Decline Parameters"]
    assert [s.layout for s in deck.slides].count("SECTION") == 1
    sections = []
    for s in deck.slides:
        if not sections or sections[-1] != s.section:
            sections.append(s.section)
    assert sections == ["Overview", "Oil", "Gas", "Water", "Appendix"]
    assert len(deck.slides) == 12
    assert titles[9:] == ["Analog Wells", "Analog Wells (1 of 2)", "Analog Wells (2 of 2)"]
    assert deck.slides[9].layout == "SECTION"
    assert deck.file_stem == "Wolfcamp_Martin_2017_21_Type_Curve_Oct2026_v1"
    assert deck.footer_left.startswith("Schaper Energy Consulting")
    assert deck.footer_right == "Production data through Sep 2025"


def test_one_number_one_value(three_phase):
    """SE QA: the P50 on the cover, the summary table and the phase slide
    are written the same way."""
    deck = R.build_deck(three_phase)
    for f in three_phase.fluids:
        p50 = f"{R.fmt_vol(f.stats['P50'])} {f.unit}"
        assert p50 in deck.slides[0].fields["stats"]
        summ = next(s for s in deck.slides if s.fields.get("title") == "EUR Summary by Phase")
        table = next(b for b in summ.blocks if b.kind == "table")
        row = next(r for r in table.data["rows"] if r[0] == f.name)
        assert row[3] == p50
        tc = next(s for s in deck.slides if s.fields.get("title") == f"{f.name} Type Curve")
        assert any(b.kind == "bullets" and b.data["heading"] == f"P50 of {p50}" for b in tc.blocks)


def test_single_phase_without_name_or_disclaimer():
    oil, _ = _fluid()
    inp = _input([oil], name="", disclaimer="")
    deck = R.build_deck(inp)
    titles = [s.fields.get("title") for s in deck.slides]
    assert titles[0] == "Type Curve Report" and "Disclaimer" not in titles
    assert not any(t and t.startswith("Gas") for t in titles)
    assert deck.file_stem.startswith("SE_Type_Curve_")


def test_text_fits_its_boxes(three_phase):
    deck = R.build_deck(three_phase)
    for s in deck.slides:
        if s.layout == "CONTENT":
            assert R.text_width(s.fields["title"], s.fields["title_size"], "head", True) <= 11.2 + 1e-6
            assert R.text_width(s.fields["sub"], s.fields["sub_size"], italic=True) <= 11.4 + 0.15
        for b in s.blocks:
            if b.kind == "bullets":
                assert R.bullets_height(b) <= b.h + 0.02, (s.fields.get("title"), b.data["heading"])
            if b.kind == "table":
                for row in b.data["rows"]:
                    for ci, cell in enumerate(row):
                        assert R.text_width(cell, b.data["size"], bold=True) <= b.data["col_w"][ci] - 0.08 + 1e-6
            if b.kind == "card":
                assert R.text_width(b.data["value"], b.data["vsize"], "head", True) <= b.w - 0.2 + 1e-6


def test_long_names_shrink_or_clip():
    oil, _ = _fluid()
    long = "Very Long Project Name For The Northern Delaware Basin Wolfcamp Development Area"
    deck = R.build_deck(_input([oil], name=long))
    cov = deck.slides[0].fields
    assert cov["title_size"] < R.SIZE["coverTitle"]
    assert R.text_width(cov["title"], cov["title_size"], "head", True) <= 11.4 + 1e-6 or cov["title_size"] == 26
    clipped = R.clip_text("A" * 200, 1.0, 9)
    assert clipped.endswith("…") and R.text_width(clipped, 9) <= 1.0


# ---------------------------------------------------------------- PPTX
def test_pptx_uses_the_se_template(rendered):
    from pptx import Presentation
    from pptx.util import Emu
    deck, pptx, _ = rendered
    prs = Presentation(pptx)
    assert round(Emu(prs.slide_width).inches, 2) == 13.33 and round(Emu(prs.slide_height).inches, 2) == 7.5
    assert len(prs.slides) == len(deck.slides)
    assert [s.slide_layout.name for s in prs.slides] == [s.layout for s in deck.slides]
    content = next(l for l in prs.slide_layouts if l.name == "CONTENT")
    footer = [s.text_frame.text for s in content.shapes if not s.is_placeholder and s.has_text_frame]
    assert deck.footer_left in footer and deck.footer_right in footer
    for i, (s, sd) in enumerate(zip(prs.slides, deck.slides), 1):
        if sd.layout == "CONTENT":
            nums = [sh.text_frame.text for sh in s.placeholders if "SLIDE_NUMBER" in str(sh.placeholder_format.type)]
            assert nums == [str(i)]
            assert s.placeholders[101].text_frame.text == sd.fields["title"]
        for ph in s.placeholders:                       # no empty prompt placeholders left
            assert ph.text_frame.text.strip() != ""
    cover = prs.slides[0]
    texts = [sh.text_frame.text for sh in cover.placeholders]
    assert deck.slides[0].fields["title"] in texts


def test_pptx_native_charts_and_sections(rendered):
    from pptx import Presentation
    _, pptx, _ = rendered
    prs = Presentation(pptx)
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    names = [c.name for c in charts]
    assert names.count("type curve chart") == 3 and names.count("cumulative chart") == 3
    assert len(names) == 6                     # two charts per phase, nothing else
    z = zipfile.ZipFile(pptx)
    xml = z.read("ppt/presentation.xml").decode()
    assert re.findall(r'<p14:section name="([^"]+)"', xml) == ["Overview", "Oil", "Gas", "Water", "Appendix"]
    tc = next(c for c in charts if c.name == "type curve chart")
    assert 'logBase val="10"' in z.read(tc.chart.part.partname.lstrip("/")).decode()
    for n in z.namelist():                              # schema order: roundedCorners after date1904
        if re.match(r"ppt/charts/chart\d+\.xml$", n):
            x = z.read(n).decode()
            if "<c:date1904" in x:
                assert x.index("<c:date1904") < x.index("<c:roundedCorners")
    for n in z.namelist():                              # unique shape ids per slide
        if re.match(r"ppt/slides/slide\d+\.xml$", n):
            ids = re.findall(r'<p:cNvPr id="(\d+)"', z.read(n).decode())
            assert len(ids) == len(set(ids)), n


def test_pptx_tables_and_bullets_follow_the_standard(rendered):
    from pptx import Presentation
    _, pptx, _ = rendered
    prs = Presentation(pptx)
    slide = next(s for s in prs.slides if s.slide_layout.name == "CONTENT"
                 and s.placeholders[101].text_frame.text == "EUR Summary by Phase")
    tbl = next(sh for sh in slide.shapes if sh.has_table).table
    head = tbl.cell(0, 0)
    assert str(head.fill.fore_color.rgb) == R.C["MAROON"]
    assert head.text_frame.paragraphs[0].runs[0].font.bold
    assert str(tbl.cell(1, 0).fill.fore_color.rgb) == R.C["CARD"]
    assert str(tbl.cell(2, 0).fill.fore_color.rgb) == R.C["WHITE"]
    bullets = next(sh for sh in slide.shapes if sh.name.startswith("panel bullets"))
    xml = bullets._element.xml
    assert 'buChar char="•"' in xml and R.C["MAROON"] in xml
    lead = bullets.text_frame.paragraphs[0].runs[0]
    assert lead.font.bold and lead.text.endswith(". ")


# ---------------------------------------------------------------- PDF
@pytest.mark.skipif(shutil.which("pdfinfo") is None, reason="poppler-utils not installed")
def test_pdf_is_the_same_deck(rendered):
    deck, _, pdf = rendered
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    assert f"Pages:           {len(deck.slides)}" in info
    assert "960 x 540 pts" in info
    for i in (1, 3, 6):
        txt = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", pdf, "-"],
                             capture_output=True, text=True).stdout
        title = deck.slides[i - 1].fields["title"]
        assert " ".join(title.split()[:2]) in txt.replace(" ", " ")
    txt3 = subprocess.run(["pdftotext", "-f", "3", "-l", "3", pdf, "-"], capture_output=True, text=True).stdout
    assert "Production data through Sep 2025" in txt3 and "Schaper Energy Consulting" in txt3


# ---------------------------------------------------------------- phase colours and cumulative
def test_phase_colours():
    assert R.phase_color("Oil") == "GREEN" and R.phase_color("gas") == "RED" and R.phase_color("Water") == "BLUE"
    styles = {n: (c, d) for n, c, d, _ in R.tc_series("GREEN")}
    assert styles["P50"] == ("GREEN", "solid") and styles["P10"][0] == "GREEN" and styles["P90"][0] == "GREEN"
    assert len({styles["P10"][1], styles["P50"][1], styles["P90"][1]}) == 3      # told apart by line style


def test_pptx_type_curves_use_the_phase_colour(rendered):
    from pptx import Presentation
    _, pptx, _ = rendered
    prs = Presentation(pptx)
    want = {"Oil": R.C["GREEN"], "Gas": R.C["RED"], "Water": R.C["BLUE"]}
    for s in prs.slides:
        if s.slide_layout.name != "CONTENT":
            continue
        title = s.placeholders[101].text_frame.text
        phase = title.split(" ")[0]
        if not title.endswith("Type Curve") or phase not in want:
            continue
        charts = {sh.name: sh.chart for sh in s.shapes if sh.has_chart}
        assert set(charts) == {"type curve chart", "cumulative chart"}
        for ch in charts.values():
            sers = {ser.name: ser for ser in ch.plots[0].series}
            assert str(sers["P50"].format.line.color.rgb) == want[phase]
            assert str(sers["P10"].format.line.color.rgb) == want[phase]
            assert str(sers["Mean (average well)"].format.line.color.rgb) == R.C["INK"]


def _colour_share(fig, hexcol):
    import io
    from PIL import Image
    im = np.asarray(Image.open(io.BytesIO(R.figure_png(fig, dpi=60))).convert("RGB")).reshape(-1, 3).astype(int)
    target = np.array([int(hexcol[i:i + 2], 16) for i in (0, 2, 4)])
    return float(np.mean(np.abs(im - target).sum(axis=1) < 40))


def test_pdf_figures_use_the_phase_colour(three_phase):
    for f in three_phase.fluids:
        col = R.phase_color(f.name)
        eurs = f.wells.loc[f.wells["in_tc"], "EUR"].to_numpy(dtype=float)
        probit = R.probit_figure(eurs, f, 7.2, 4.8, col)
        assert _colour_share(probit, R.C[col]) > 0.0005
        assert _colour_share(probit, R.C["MAROON"]) < 0.0001
        spec = dict(curves=f.curves, lines=f.lines, name=f.name, n_wells=f.n_tc, color=col, unit=f.unit, title="t")
        assert _colour_share(R.tc_figure(spec, 4.0, 5.3), R.C[col]) > 0.001
        assert _colour_share(R.cum_figure(spec, 4.0, 5.3), R.C[col]) > 0.001


def test_cumulative_ends_at_the_eur(three_phase):
    for f in three_phase.fluids:
        cums, hist = R.cum_series(dict(curves=f.curves, lines=f.lines))
        for k in ("P90", "P50", "P10", "Mean"):
            assert cums[k][1][-1] == pytest.approx(f.stats[k], rel=1e-9)
            assert np.all(np.diff(cums[k][1]) >= 0)
        assert len(hist) == len(f.lines)
        assert R._p50_cum(f, 12) == pytest.approx(f.curves["P50"].to_numpy()[:12].sum() / 1000.0)
        assert R._p50_cum(f, 600) == pytest.approx(f.stats["P50"], rel=1e-9)


def test_pptx_disclaimer_has_no_subtitle(rendered):
    from pptx import Presentation
    _, pptx, _ = rendered
    s = Presentation(pptx).slides[1]
    assert s.placeholders[101].text_frame.text == "Disclaimer"
    assert 102 not in [ph.placeholder_format.idx for ph in s.placeholders]
    assert "Read this before" not in "".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)


def test_two_slides_per_phase_with_b_summary(three_phase):
    deck = R.build_deck(three_phase)
    for f in three_phase.fluids:
        assert sum(1 for sd in deck.slides if sd.section == f.name) == 2
        dec = next(sd for sd in deck.slides if sd.fields.get("title") == f"{f.name} EUR Distribution and Decline Parameters")
        caps = [b for b in dec.blocks if b.kind == "text" and b.data["name"] == "caption"]
        text = " ".join("".join(t for t, _ in p) for b in caps for p in b.data["paras"]).replace("\u00a0", " ")
        assert f"b of the {len(f.b_values)} wells: median" in text
        assert all(b.y + b.h <= 6.9 for b in dec.blocks)          # nothing runs onto the footer
